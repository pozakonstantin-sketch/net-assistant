"""Присланные пользователем файлы: извлечение текста, хранение, чтение для ИИ.

Оригинал сохраняется на SSD роутера (см. :mod:`core.filestore`) и в локальный
кэш ``data/files``. Из текста делается sidecar-файл, чтобы планировщик мог
читать содержимое без повторного разбора Word/Excel/PDF.

Поддерживаются: docx, xlsx, pdf, doc, xls, ppt, rtf, odt/ods/odp, txt, csv,
md, json, html, xml, логи, скрипты, zip (список) и обычный текст.
"""

from __future__ import annotations

import json
import logging
import os
import posixpath
import re
import secrets
import subprocess
import threading
import time
import zipfile
from html import unescape
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from . import filestore
from .config import DATA_DIR, settings

log = logging.getLogger("documents")

FILES_DIR = DATA_DIR / "files"
INDEX = FILES_DIR / "index.jsonl"
MAX_TEXT_CHARS = 400_000
PROMPT_CHARS = 12_000
READ_DEFAULT = 8_000
SEARCH_SCAN_CHARS = 2_000_000

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".log",
    ".ini", ".cfg", ".conf", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".sh", ".ps1", ".bat", ".sql", ".srt", ".vtt",
    ".toml", ".env", ".properties",
}

KIND_LABELS = {
    ".docx": "документ Word",
    ".doc": "документ Word (старый формат)",
    ".xlsx": "таблица Excel",
    ".xls": "таблица Excel (старый формат)",
    ".pdf": "PDF",
    ".pptx": "презентация PowerPoint",
    ".ppt": "презентация PowerPoint (старый формат)",
    ".rtf": "документ RTF",
    ".odt": "документ OpenDocument",
    ".ods": "таблица OpenDocument",
    ".odp": "презентация OpenDocument",
    ".zip": "архив ZIP",
}

_lock = threading.RLock()


class FileError(ValueError):
    """Ошибка работы с присланным файлом."""


class ExtractionError(RuntimeError):
    """Не удалось извлечь текст из файла."""


# ─────────────────────────────── утилиты ───────────────────────────────


def safe_name(name: str) -> str:
    """Имя без каталогов и спецсимволов, с сохранением расширения."""
    base = Path(str(name or "file")).name
    base = re.sub(r"[^\w.\- ()\[\]]+", "_", base, flags=re.UNICODE).strip(" .")
    if not base:
        base = "file"
    if len(base) > 120:
        stem, suffix = Path(base).stem[:100], Path(base).suffix[:19]
        base = stem + suffix
    return base


def _size_text(size: int) -> str:
    for unit, scale in (("ГБ", 1024 ** 3), ("МБ", 1024 ** 2), ("КБ", 1024)):
        if size >= scale:
            return f"{size / scale:.1f} {unit}"
    return f"{size} Б"


def _clip(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[:limit] + "…"


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1251", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", "replace")


def _run(command: list[str], timeout: int = 60) -> str:
    try:
        proc = subprocess.run(
            command, capture_output=True, timeout=timeout,
            env={**os.environ, "LANG": "C.UTF-8", "LC_ALL": "C.UTF-8"},
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise ExtractionError(str(exc)) from exc
    if proc.returncode != 0:
        detail = proc.stderr.decode("utf-8", "replace").strip()[:300]
        raise ExtractionError(detail or f"код выхода {proc.returncode}")
    return proc.stdout.decode("utf-8", "replace")


def _strip_html(text: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
    text = re.sub(r"(?i)<br\s*/?>", "\n", text)
    text = re.sub(r"(?i)</(p|div|tr|li|h[1-6])>", "\n", text)
    text = re.sub(r"(?i)</t[dh]>", "\t", text)
    text = re.sub(r"<[^>]+>", " ", text)
    return _clean(unescape(text))


# ─────────────────────────── извлечение текста ───────────────────────────


def _extract_docx(path: Path) -> tuple[str, dict]:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        line = paragraph.text.strip()
        if line:
            parts.append(line)
    for number, table in enumerate(document.tables, 1):
        parts.append(f"\nТаблица {number}:")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    meta = {"paragraphs": len(document.paragraphs), "tables": len(document.tables)}
    return _clean("\n".join(parts)), meta


def _extract_xlsx(path: Path) -> tuple[str, dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    total_rows = 0
    sheet_count = len(workbook.sheetnames)
    try:
        for sheet in workbook.worksheets:
            parts.append(f"\n=== Лист «{sheet.title}» ===")
            rows = 0
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if not any(value.strip() for value in values):
                    continue
                parts.append("\t".join(values).rstrip())
                rows += 1
                total_rows += 1
                if rows >= 3000 or sum(len(p) for p in parts) > MAX_TEXT_CHARS:
                    parts.append(f"… (лист «{sheet.title}» обрезан)")
                    break
            if sum(len(p) for p in parts) > MAX_TEXT_CHARS:
                break
    finally:
        workbook.close()
    return _clean("\n".join(parts)), {"sheets": sheet_count, "rows": total_rows}


def _extract_pdf(path: Path) -> tuple[str, dict]:
    try:
        text = _run(["pdftotext", "-layout", "-enc", "UTF-8", str(path), "-"])
    except ExtractionError:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    meta: dict[str, Any] = {}
    try:
        info = _run(["pdfinfo", str(path)])
        for line in info.splitlines():
            if line.lower().startswith("pages:"):
                meta["pages"] = int(line.split(":", 1)[1].strip())
    except (ExtractionError, ValueError):
        pass
    return _clean(text), meta


def _extract_odf(path: Path) -> tuple[str, dict]:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("content.xml").decode("utf-8", "replace")
    xml = xml.replace("</text:p>", "\n").replace("</text:h>", "\n")
    xml = xml.replace("<text:tab/>", "\t").replace("<text:line-break/>", "\n")
    text = re.sub(r"<[^>]+>", "", unescape(xml))
    return _clean(text), {}


def _extract_zip(path: Path) -> tuple[str, dict]:
    lines: list[str] = []
    with zipfile.ZipFile(path) as archive:
        items = archive.infolist()
        for item in items[:200]:
            lines.append(f"{item.filename} ({_size_text(item.file_size)})")
        if len(items) > 200:
            lines.append(f"… всего файлов: {len(items)}")
    return "\n".join(lines), {"entries": len(items)}


def extract(path: Path) -> dict:
    """Вернуть ``{"text", "kind", "meta"}``; ошибки не выбрасываются наружу."""
    suffix = path.suffix.lower()
    kind = KIND_LABELS.get(suffix, suffix.lstrip(".").upper() or "файл")
    try:
        if suffix == ".docx":
            text, meta = _extract_docx(path)
        elif suffix == ".xlsx":
            text, meta = _extract_xlsx(path)
        elif suffix == ".pdf":
            text, meta = _extract_pdf(path)
        elif suffix in {".odt", ".ods", ".odp"}:
            text, meta = _extract_odf(path)
        elif suffix == ".zip":
            text, meta = _extract_zip(path)
        elif suffix == ".doc":
            try:
                text, meta = _run(["catdoc", "-d", "utf-8", "-w", str(path)]), {}
            except ExtractionError:
                text, meta = _run(["antiword", "-w", "0", str(path)]), {}
        elif suffix == ".xls":
            text, meta = _run(["xls2csv", "-d", "utf-8", str(path)]), {}
        elif suffix == ".ppt":
            text, meta = _run(["catppt", "-d", "utf-8", str(path)]), {}
        elif suffix == ".pptx":
            with zipfile.ZipFile(path) as archive:
                slides = sorted(
                    (name for name in archive.namelist()
                     if re.match(r"ppt/slides/slide\d+\.xml$", name)),
                    key=lambda name: int(re.search(r"(\d+)", name).group(1)),
                )
                chunks = []
                for name in slides:
                    xml = archive.read(name).decode("utf-8", "replace")
                    chunks.append(_clean(re.sub(r"<[^>]+>", " ", xml)))
            text, meta = "\n\n".join(chunks), {"slides": len(slides)}
        elif suffix == ".rtf":
            text, meta = _run(["unrtf", "--text", "--nopict", str(path)]), {}
        elif suffix in TEXT_SUFFIXES:
            text = _decode(path.read_bytes())
            if suffix in {".html", ".htm"}:
                text = _strip_html(text)
            meta = {}
        else:
            raw = path.read_bytes()
            if b"\x00" not in raw[:4096]:
                text, meta = _decode(raw), {}
            else:
                described = _run(["file", "--brief", str(path)]).strip()
                text = f"(двоичный файл: {described}; текст не извлекался)"
                meta = {"binary": True}
        text = _clip(_clean(text), MAX_TEXT_CHARS)
        return {"text": text, "kind": kind, "meta": meta, "error": ""}
    except Exception as exc:  # noqa: BLE001 - файл не должен ронять ядро
        log.warning("не удалось извлечь текст из %s: %s", path.name, exc)
        return {"text": "", "kind": kind, "meta": {}, "error": str(exc)}


# ─────────────────────────────── хранилище ───────────────────────────────


def _append(entry: dict) -> None:
    FILES_DIR.mkdir(parents=True, exist_ok=True)
    with _lock:
        with INDEX.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(entry, ensure_ascii=False) + "\n")


def entries() -> list[dict]:
    if not INDEX.exists():
        return []
    result: list[dict] = []
    try:
        for line in INDEX.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line:
                continue
            try:
                result.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    except OSError as exc:
        log.warning("не удалось прочитать индекс файлов: %s", exc)
    return result


def find(entry_id: str) -> dict:
    wanted = (entry_id or "").strip()
    if not wanted:
        raise FileError("укажите id файла (см. file.list)")
    items = entries()
    for item in reversed(items):
        if item.get("id") == wanted:
            return item
    matches = [item for item in items if str(item.get("id", "")).startswith(wanted)]
    if len(matches) == 1:
        return matches[0]
    by_name = [item for item in items if str(item.get("name", "")).lower() == wanted.lower()]
    if len(by_name) == 1:
        return by_name[0]
    if len(matches) > 1 or len(by_name) > 1:
        ids = ", ".join(str(item.get("id")) for item in (matches or by_name)[:10])
        raise FileError(f"несколько файлов подходят под «{wanted}»: {ids}")
    raise FileError(f"файл «{wanted}» не найден; список — file.list")


def save(data: bytes, filename: str, session: str, source: str, caption: str = "") -> dict:
    """Сохранить присланный файл локально и на SSD роутера, извлечь текст."""
    limit = settings.files_max_mb * 1024 * 1024
    if len(data) > limit:
        raise FileError(f"файл больше {settings.files_max_mb} МБ")
    if not data:
        raise FileError("пустой файл")

    FILES_DIR.mkdir(parents=True, exist_ok=True)
    name = safe_name(filename)
    entry_id = f"{time.strftime('%Y%m%d_%H%M%S')}_{secrets.token_hex(3)}"
    local = FILES_DIR / f"{entry_id}__{name}"
    local.write_bytes(data)

    info = extract(local)
    text_path = FILES_DIR / f"{entry_id}.txt"
    text_path.write_text(info["text"], encoding="utf-8")

    remote_dir = time.strftime("%Y-%m-%d")
    router_path = ""
    router_error = ""
    try:
        uploaded = filestore.upload_many(
            [(local, f"{entry_id}__{name}"), (text_path, f"{entry_id}.txt")],
            subdir=remote_dir,
        )
        router_path = uploaded[0] if uploaded else ""
    except filestore.StorageError as exc:
        router_error = str(exc)
        log.warning("файл остался только на ВМ: %s", exc)

    entry = {
        "id": entry_id,
        "name": name,
        "kind": info["kind"],
        "size": len(data),
        "size_text": _size_text(len(data)),
        "chars": len(info["text"]),
        "meta": info["meta"],
        "extract_error": info["error"],
        "uploaded": time.time(),
        "session": session,
        "source": source,
        "caption": caption[:1000],
        "local": str(local),
        "text_path": str(text_path),
        "router_path": router_path,
        "router_error": router_error,
        "preview": _clip(re.sub(r"\s+", " ", info["text"]).strip(), 300),
    }
    _append(entry)
    log.info("принят файл %s (%s, %s), роутер: %s",
             name, info["kind"], entry["size_text"], router_path or "нет")
    return entry


# ─────────────────────────────── для ИИ ───────────────────────────────


def chat_block(entry: dict, caption: str = "") -> str:
    """Текст реплики пользователя с содержимым файла для планировщика."""
    lines = []
    header = f"[Пользователь прислал файл «{entry['name']}»"
    if entry.get("kind"):
        header += f" ({entry['kind']}"
        if entry.get("meta", {}).get("pages"):
            header += f", {entry['meta']['pages']} стр."
        if entry.get("meta", {}).get("sheets"):
            header += f", листов: {entry['meta']['sheets']}"
        header += ")"
    header += f", {entry.get('size_text') or _size_text(int(entry.get('size') or 0))}."
    header += f" ID файла: {entry['id']}."
    if entry.get("router_path"):
        header += f" Оригинал сохранён на SSD роутера: {entry['router_path']}."
    else:
        header += (f" Оригинал сохранён только на ВМ ({entry.get('local')}); "
                   f"SSD роутера недоступен: {entry.get('router_error') or 'нет связи'}.")
    lines.append(header)

    text = ""
    text_path = entry.get("text_path")
    if text_path and Path(text_path).is_file():
        try:
            text = Path(text_path).read_text(encoding="utf-8")
        except OSError:
            text = ""
    if text.strip():
        lines.append(f"Извлечённый текст ({len(text)} знаков), фрагмент:")
        lines.append(_clip(text, PROMPT_CHARS))
        if len(text) > PROMPT_CHARS:
            lines.append("… (фрагмент обрезан; полный текст — core::file.read)")
    elif entry.get("extract_error"):
        lines.append(f"Текст извлечь не удалось: {entry['extract_error']}")
    else:
        lines.append("Текст в файле не найден (возможно, скан или картинка).")
    lines.append("]")

    block = "\n".join(lines)
    caption = (caption or "").strip()
    return f"{caption}\n\n{block}" if caption else block


def display_text(entry: dict, caption: str = "") -> str:
    """Короткая реплика для истории и панели (без килобайтов текста файла)."""
    parts = [entry.get("kind"), entry.get("size_text")]
    head = f"📎 Файл «{entry['name']}»"
    details = [str(part) for part in parts if part]
    if details:
        head += " — " + ", ".join(details)
    if entry.get("router_path"):
        head += f"\n💾 SSD роутера: {entry['router_path']}"
    elif entry.get("router_error"):
        head += f"\n💾 Сохранён на ВМ: SSD роутера недоступен ({entry['router_error']})"
    caption = (caption or "").strip()
    return f"{head}\n{caption}" if caption else head


def prompt_hint(limit: int = 5) -> str:
    """Короткая сводка о файлах для системного промпта."""
    items = entries()
    base = settings.router_files_dir
    if not items:
        return (f"Присланных файлов пока нет. Каталог для них на SSD роутера: {base}. "
                f"Когда пользователь пришлёт файл, используй core::file.list/read/search.")
    lines = [
        f"Присланные файлы: {len(items)} шт., оригиналы на SSD роутера в {base}.",
        "Недавние:",
    ]
    for item in items[-limit:][::-1]:
        lines.append(f"- {item.get('id')} — {item.get('name')} "
                     f"({item.get('kind')}, {item.get('size_text')}, знаков: {item.get('chars')})")
    lines.append("Читай содержимое через core::file.read(id=...), ищи через core::file.search(query=...).")
    return "\n".join(lines)


def core_action(action: str, args: dict) -> dict:
    """Локальные действия ядра: file.list / file.read / file.search."""
    if action == "file.list":
        limit = max(1, min(int(args.get("limit") or 15), 100))
        items = entries()[-limit:][::-1]
        return {
            "count": len(entries()),
            "files": [
                {
                    "id": item.get("id"),
                    "name": item.get("name"),
                    "kind": item.get("kind"),
                    "size": item.get("size_text"),
                    "chars": item.get("chars"),
                    "uploaded": time.strftime("%Y-%m-%d %H:%M",
                                              time.localtime(item.get("uploaded") or 0)),
                    "router_path": item.get("router_path") or "",
                    "preview": item.get("preview") or "",
                }
                for item in items
            ],
        }

    if action == "file.read":
        item = find(str(args.get("id") or ""))
        offset = max(0, int(args.get("offset") or 0))
        limit = max(1, min(int(args.get("limit") or READ_DEFAULT), 20000))
        text_path = Path(str(item.get("text_path") or ""))
        if not text_path.is_file():
            raise FileError("текст файла не сохранён; возможно, формат не поддерживается")
        text = text_path.read_text(encoding="utf-8")
        chunk = text[offset:offset + limit]
        return {
            "id": item.get("id"),
            "name": item.get("name"),
            "kind": item.get("kind"),
            "router_path": item.get("router_path") or "",
            "chars": len(text),
            "offset": offset,
            "text": chunk,
            "next_offset": offset + len(chunk) if offset + len(chunk) < len(text) else None,
            "extract_error": item.get("extract_error") or "",
        }

    if action == "file.search":
        query = str(args.get("query") or "").strip()
        if not query:
            raise FileError("пустой запрос поиска")
        limit = max(1, min(int(args.get("limit") or 10), 50))
        needle = query.lower()
        found: list[dict] = []
        for item in reversed(entries()):
            text_path = Path(str(item.get("text_path") or ""))
            if not text_path.is_file():
                continue
            try:
                text = text_path.read_text(encoding="utf-8")[:SEARCH_SCAN_CHARS]
            except OSError:
                continue
            position = text.lower().find(needle)
            if position < 0:
                continue
            start = max(0, position - 160)
            snippet = text[start:position + len(query) + 160].replace("\n", " ")
            found.append({
                "id": item.get("id"),
                "name": item.get("name"),
                "router_path": item.get("router_path") or "",
                "snippet": _clip(snippet, 400),
            })
            if len(found) >= limit:
                break
        return {"query": query, "count": len(found), "matches": found}

    raise FileError(f"неизвестное действие с файлами: {action}")


def storage_info() -> dict:
    """Сводка для /health и панели."""
    items = entries()
    return {
        "count": len(items),
        "chars": sum(int(item.get("chars") or 0) for item in items),
        "remote_dir": settings.router_files_dir,
        "router": filestore.enabled(),
        "last": items[-1].get("name") if items else "",
    }
